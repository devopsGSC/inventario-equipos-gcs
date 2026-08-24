import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "Equipos"

navy  = "0F1724"
blue  = "3B82F6"
white = "FFFFFF"
amber = "FEF3C7"
thin  = Side(style='thin', color="CBD5E1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Fila 1: Título
ws.merge_cells("A1:N1")
ws["A1"] = "InventarioTI — Plantilla de Carga Masiva de Equipos"
ws["A1"].font = Font(name="Arial", bold=True, size=13, color=white)
ws["A1"].fill = PatternFill("solid", fgColor=navy)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 28

# Fila 2: Instrucción
ws.merge_cells("A2:N2")
ws["A2"] = "Complete los campos desde la fila 5. Campos con * son obligatorios. No modifique los encabezados. Descargue tipos válidos en la hoja 'Referencia'."
ws["A2"].font = Font(name="Arial", size=9, color="92400E")
ws["A2"].fill = PatternFill("solid", fgColor=amber)
ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 20

# Fila 3: Separador
ws.row_dimensions[3].height = 6

# Fila 4: Encabezados
headers = [
    ("NombreEquipo *",  "A", 28),
    ("TipoEquipo *",    "B", 18),
    ("Marca *",         "C", 18),
    ("Modelo *",        "D", 22),
    ("NumeroSerie *",   "E", 22),
    ("IMEI",            "F", 20),
    ("Accesorios",      "G", 28),
    ("Costo",           "H", 12),
    ("FechaCompra",     "I", 16),
    ("FechaGarantia",   "J", 16),
]

for i, (title, col, width) in enumerate(headers, 1):
    cell = ws.cell(row=4, column=i, value=title)
    cell.font      = Font(name="Arial", bold=True, size=10, color=white)
    cell.fill      = PatternFill("solid", fgColor=blue)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = border
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[4].height = 22

# Filas 5-24: Datos con formato zebra
for row in range(5, 25):
    bg = "FFFFFF" if row % 2 == 0 else "F8FAFC"
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill      = PatternFill("solid", fgColor=bg)
        cell.border    = border
        cell.font      = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="center")
    # Formato fecha columnas I y J
    ws.cell(row=row, column=9).number_format  = "DD/MM/YYYY"
    ws.cell(row=row, column=10).number_format = "DD/MM/YYYY"
    # Formato número columna H
    ws.cell(row=row, column=8).number_format  = "#,##0.00"
    ws.row_dimensions[row].height = 18

# Validación desplegable para TipoEquipo (columna B)
dv_tipo = DataValidation(
    type="list",
    formula1='"Laptop,Celular,Tablet,Otro"',
    allow_blank=True,
    showErrorMessage=True,
    errorTitle="Tipo inválido",
    error="Valor no válido. Use: Laptop, Celular, Tablet u Otro"
)
ws.add_data_validation(dv_tipo)
dv_tipo.sqref = "B5:B24"

# Hoja Referencia
ref = wb.create_sheet("Referencia")
ref["A1"] = "Tipos de Equipo válidos"
ref["A1"].font = Font(name="Arial", bold=True, size=11, color=white)
ref["A1"].fill = PatternFill("solid", fgColor=navy)
ref.column_dimensions["A"].width = 25
ref.column_dimensions["B"].width = 35

tipos = ["Laptop", "Celular", "Tablet", "Otro"]
for i, t in enumerate(tipos, 2):
    ref.cell(row=i, column=1, value=t).font = Font(name="Arial", size=10)
    ref.cell(row=i, column=1).border = border

ref["A8"] = "Formato fechas:"
ref["A8"].font = Font(name="Arial", bold=True, size=10)
ref["A9"] = "DD/MM/YYYY"
ref["A9"].font = Font(name="Arial", size=10)
ref["A10"] = "Ej: 15/06/2025"
ref["A10"].font = Font(name="Arial", size=10, color="6B7280")

ref["A12"] = "Accesorios:"
ref["A12"].font = Font(name="Arial", bold=True, size=10)
ref["A13"] = "Separados por coma"
ref["A13"].font = Font(name="Arial", size=10)
ref["A14"] = "Ej: Cargador, Funda"
ref["A14"].font = Font(name="Arial", size=10, color="6B7280")

# Guardar en wwwroot del proyecto
import os
ruta = os.path.join(os.path.dirname(__file__), "wwwroot", "plantilla_equipos.xlsx")
wb.save(ruta)
print(f"Plantilla generada en: {ruta}")
